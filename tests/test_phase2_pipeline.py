"""Phase 2 acceptance: synthetic records only, no live requests or original edits."""
import csv
from hashlib import sha256
import io
import json
from pathlib import Path
import sqlite3
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from jobhunter_app.store import Store, SCHEMA, DISCOVERY_SCHEMA
from jobhunter_app.migration import migrate, history, restore_backup
from jobhunter_app.intake import parse_jobs, export_jobs
from jobhunter_app.pipeline import import_records, export_history, export_uaa
from jobhunter_app.files import publish_new
from tests.test_local_store import JOB, PROFILE

@pytest.fixture
def store(tmp_path):
    s=Store(tmp_path/'app.db');s.save_profile(PROFILE);return s

def legacy(tmp_path, records=None):
    d=tmp_path/'legacy';d.mkdir();(d/'evaluations.json').write_text(json.dumps(records if records is not None else [JOB]));return d

def counts(s):
    with s.connect() as db:
        return {t:db.execute('SELECT COUNT(*) FROM '+t).fetchone()[0] for t in ('jobs','historical_records','migration_batches','applications','status_history','audit')}

def xlsx(rows, headers=None):
    import openpyxl
    book=openpyxl.Workbook();s=book.active
    if headers:s.append(headers)
    for row in rows:s.append(row)
    b=io.BytesIO();book.save(b);book.close();return b.getvalue()

@pytest.mark.parametrize('version',[1,2])
def test_schema_upgrade_snapshot_and_idempotence(tmp_path,version):
    p=tmp_path/'old.db'
    with sqlite3.connect(p) as db:
        db.executescript(SCHEMA+(DISCOVERY_SCHEMA if version==2 else '')+f'PRAGMA user_version={version};')
        db.execute('INSERT INTO profiles VALUES(?,?,?)',('default',json.dumps(PROFILE),'before'))
    s=Store(p)
    assert s.profile()==PROFILE
    with s.connect() as db:
        assert db.execute('PRAGMA user_version').fetchone()[0]==5
        assert db.execute('PRAGMA foreign_key_check').fetchall()==[]
    snapshots=list(tmp_path.glob('*.bak'));assert len(snapshots)==1
    with sqlite3.connect(snapshots[0]) as db:assert db.execute('PRAGMA user_version').fetchone()[0]==version
    Store(p);assert len(list(tmp_path.glob('*.bak')))==1

def test_migration_preview_apply_rerun_is_read_only(store,tmp_path):
    d=legacy(tmp_path);original=(d/'evaluations.json').read_bytes();before=counts(store)
    p=migrate(store,d);assert p['preview']['inserted']==1 and p['new_history_records']==1
    assert counts(store)==before and not (d/'migration-backups').exists()
    applied=migrate(store,d,False);after=counts(store)
    assert Path(applied['backup'],'evaluations.json').read_bytes()==original
    again=migrate(store,d,False);assert again['already_applied'] and counts(store)==after
    assert (d/'evaluations.json').read_bytes()==original and history(store)[0]['payload']==JOB

def test_incremental_migration_records_new_observations(store,tmp_path):
    d=legacy(tmp_path);migrate(store,d,False)
    (d/'evaluations.json').write_text(json.dumps([{**JOB,'global_score':4,'evaluated_at':'2026-10-29'}]))
    migrate(store,d,False)
    assert len(history(store))==2 and store.summary()['total_jobs']==1
    assert store.summary()['matched_jobs']==0  # History scores do not become current matches.

@pytest.mark.parametrize('kind',['recommendations','pipeline','queue','scan'])
def test_legacy_states_are_not_application_facts(store,tmp_path,kind):
    d=tmp_path/'legacy';d.mkdir()
    if kind=='recommendations':
        (d/'applications.md').write_text('| # | Date | Company | Role | Score | Status | Report |\n|---|---|---|---|---|---|---|\n|1|2026-10-29|Example|Engineer|5|apply|report.md|\n')
    elif kind=='pipeline':
        (d/'pipeline.md').write_text(f"- [x] {JOB['url']} | {JOB['company']} | {JOB['title']} [Cairo]\n")
    elif kind=='queue':(d/'application_queue.jsonl').write_text(json.dumps({**JOB,'status':'ready_to_apply'})+'\n')
    else:(d/'scan-history.tsv').write_text('url\ttitle\tcompany\tstatus\n'+f"{JOB['url']}\t{JOB['title']}\t{JOB['company']}\tskipped_title\n")
    migrate(store,d,False)
    assert len(history(store))==1 and counts(store)['applications']==0 and counts(store)['status_history']==0
    if kind in ('recommendations','scan'):assert store.summary()['total_jobs']==0

def test_migration_profiles_are_isolated(store,tmp_path):
    d=legacy(tmp_path);store.save_profile(PROFILE,'other');migrate(store,d,False,'other')
    assert history(store)==[] and len(history(store,'other'))==1
    assert json.loads(export_history(store))['historical_records']==[]
    migrate(store,d,False)
    assert len(history(store))==1 and store.summary()['total_jobs']==1

@pytest.mark.parametrize('field',['company','title'])
def test_migration_conflicts_never_overwrite(store,tmp_path,field):
    store.ingest([JOB]);d=legacy(tmp_path,[{**JOB,field:'Conflicting identity'}]);before=counts(store)
    assert migrate(store,d)['conflicts'][0]['fields']==[field]
    with pytest.raises(ValueError):migrate(store,d,False)
    assert counts(store)==before and store.jobs()[0][field]==JOB[field]
    r=migrate(store,d,False,allow_partial=True)
    assert r['conflicts'] and history(store)[0]['job_id'] is None
    assert store.jobs()[0][field]==JOB[field]

def test_existing_rich_job_and_notes_remain_authoritative(store,tmp_path):
    store.ingest([{**JOB,'description':'Rich authoritative description'}]);jid=store.jobs()[0]['id'];store.set_status(jid,'saved','keep')
    d=legacy(tmp_path,[{**JOB,'description':'Short legacy text'}]);migrate(store,d,False)
    assert store.jobs()[0]['description']=='Rich authoritative description'
    assert store.jobs()[0]['notes']=='keep' and store.jobs()[0]['status']=='saved'

@pytest.mark.parametrize('failure',['ingest','history','batch'])
def test_apply_transaction_rolls_back_all_tables(store,tmp_path,failure):
    d=legacy(tmp_path);table={'ingest':'sightings','history':'historical_records','batch':'migration_batches'}[failure]
    with store.connect() as db:db.execute(f"CREATE TRIGGER fail BEFORE INSERT ON {table} BEGIN SELECT RAISE(ABORT,'fixture'); END")
    before=counts(store)
    with pytest.raises(sqlite3.DatabaseError):migrate(store,d,False)
    assert counts(store)==before
    backups=list((d/'migration-backups').glob('*/before.sqlite3'));assert len(backups)==1
    with sqlite3.connect(backups[0]) as db:assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'

def test_changed_input_between_snapshot_and_apply_fails(store,tmp_path,monkeypatch):
    d=legacy(tmp_path);original_backup=store.backup
    def modify(path):
        result=original_backup(path);(d/'evaluations.json').write_text('[]');return result
    monkeypatch.setattr(store,'backup',modify)
    before=counts(store)
    with pytest.raises(ValueError,match='input changed'):migrate(store,d,False)
    assert counts(store)==before

@pytest.mark.parametrize('name,content',[
 ('evaluations.json','[NaN]'),('evaluations.json','[{"url":"a","url":"b"}]'),
 ('jd_cache.json','[]'),('pipeline.md','- [x] malformed'),
 ('scan-history.tsv','url\ttitle\nhttps://example.org\n'),('application_queue.jsonl','not json'),
 ('applications.md','|a|b|\n|---|---|\n|1|2|3|\n')])
def test_malformed_legacy_inputs_are_reported(store,tmp_path,name,content):
    d=tmp_path/'legacy';d.mkdir();(d/name).write_text(content);before=counts(store)
    assert migrate(store,d)['file_errors']
    with pytest.raises(ValueError):migrate(store,d,False)
    assert counts(store)==before

def test_symlink_and_unknown_file_safety(store,tmp_path):
    d=tmp_path/'legacy';d.mkdir();outside=tmp_path/'outside';outside.write_text('[]');(d/'evaluations.json').symlink_to(outside)
    (d/'private.docx').write_bytes(b'not opened')
    p=migrate(store,d);assert p['file_errors'] and 'private.docx' in p['ignored_files']
    assert outside.read_text()=='[]'

def test_legacy_excel_preserves_typed_evaluation(store,tmp_path):
    d=tmp_path/'legacy';d.mkdir();content=xlsx([['Example','Engineer','https://example.org/jobs/1',4.5,'apply']],['Company','Title','URL','Score','Recommendation'])
    (d/'jobs_database.xlsx').write_bytes(content)
    migrate(store,d,False)
    assert history(store)[0]['payload']['Score']==4.5 and store.summary()['matched_jobs']==0
    assert (d/'jobs_database.xlsx').read_bytes()==content

def test_restore_new_file_no_live_clobber(store,tmp_path):
    d=legacy(tmp_path);r=migrate(store,d,False);target=tmp_path/'restored.db'
    restored=restore_backup(Path(r['backup'])/'before.sqlite3',target)
    assert restored['schema']==5 and Store(target).summary()['total_jobs']==0
    assert store.summary()['total_jobs']==1
    with pytest.raises(ValueError):restore_backup(Path(r['backup'])/'before.sqlite3',store.path)

@pytest.mark.parametrize('contents',[b'not sqlite',b''])
def test_invalid_restore_does_not_publish(store,tmp_path,contents):
    p=tmp_path/'invalid';p.write_bytes(contents);target=tmp_path/'new.db'
    with pytest.raises((ValueError,sqlite3.DatabaseError)):restore_backup(p,target)
    assert not target.exists() and p.read_bytes()==contents

def test_canonical_state_roundtrip_and_profile_scope(store,tmp_path):
    store.ingest([JOB]);jid=store.jobs()[0]['id'];store.set_status(jid,'applied','User reported','2026-12-01')
    rows=json.loads(export_jobs(store));other=Store(tmp_path/'other.db');other.save_profile(PROFILE,'chosen')
    before=counts(other);p=import_records(other,rows,'chosen',True,True);assert p['state_rows']==1 and counts(other)==before
    import_records(other,rows,'chosen',with_state=True)
    assert other.jobs(profile_id='chosen')[0]['status']=='applied' and other.jobs()[0]['status']=='discovered'
    import_records(other,rows,'chosen',with_state=True);assert counts(other)['status_history']==1
    assert other.jobs(profile_id='chosen')[0]['notes']=='User reported'

def test_state_ignored_without_explicit_flag(store):
    import_records(store,[{**JOB,'status':'applied','notes':'legacy?'}])
    assert store.jobs()[0]['status']=='discovered' and counts(store)['status_history']==0

@pytest.mark.parametrize('change',[{'status':'ready_to_apply'},{'deadline':'tomorrow'},{'notes':False},{'title':'Conflicting'}])
def test_invalid_or_conflicting_import_is_atomic(store,change):
    store.ingest([JOB]);before=counts(store);rows=[{**JOB,**change}]
    p=import_records(store,rows,dry_run=True,with_state=True);assert p['invalid'] or p['conflicts']
    with pytest.raises(ValueError):import_records(store,rows,with_state=True)
    assert counts(store)==before

def test_current_status_never_silently_overwritten(store):
    import_records(store,[{**JOB,'status':'saved'}],with_state=True);before=counts(store)
    with pytest.raises(ValueError):import_records(store,[{**JOB,'status':'applied'}],with_state=True)
    assert counts(store)==before and store.jobs()[0]['status']=='saved'

@pytest.mark.parametrize('fmt,content',[('csv',b'url,title,title\na,b,c'),('csv',b'url,title,company\na,b'),('csv',b'url,title,company\na,b,c,d'),('json',b'[NaN]'),('json',b'[{"a":1,"a":2}]')])
def test_ambiguous_file_shapes_rejected(fmt,content):
    with pytest.raises(ValueError):parse_jobs(content,fmt)

def test_excel_literal_formula_roundtrip(store):
    import_records(store,[{**JOB,'title':'=1+1','company':'@Example','description':'-literal'}])
    content=export_jobs(store,'xlsx');rows=parse_jobs(content,'xlsx')
    assert rows[0]['title']=='=1+1' and rows[0]['company']=='@Example'
    import openpyxl
    book=openpyxl.load_workbook(io.BytesIO(content),data_only=False)
    assert all(c.data_type!='f' for row in book.active for c in row);book.close()

@pytest.mark.parametrize('mode',['formula','duplicate','multiple','error','numeric'])
def test_excel_unsafe_or_ambiguous_shapes_rejected(mode):
    import openpyxl
    b=openpyxl.Workbook();s=b.active;s.append(['url','title','company'])
    s.append([JOB['url'],JOB['title'],JOB['company']])
    if mode=='formula':s['B2']='=WEBSERVICE("https://example.org")'
    elif mode=='duplicate':s['B1']='url'
    elif mode=='multiple':b.create_sheet('Hidden input')
    elif mode=='error':s['B2']='#REF!'
    else:s['B2']=123
    out=io.BytesIO();b.save(out);b.close()
    with pytest.raises(ValueError):parse_jobs(out.getvalue(),'xlsx')

def test_excel_external_relationship_rejected():
    content=xlsx([[JOB['url'],JOB['title'],JOB['company']]],['url','title','company']);out=io.BytesIO()
    with ZipFile(io.BytesIO(content)) as src,ZipFile(out,'w',ZIP_DEFLATED) as z:
        for info in src.infolist():z.writestr(info,src.read(info.filename))
        z.writestr('xl/worksheets/_rels/sheet1.xml.rels','<Relationships><Relationship TargetMode="External" Target="https://example.org"/></Relationships>')
    with pytest.raises(ValueError,match='External'):parse_jobs(out.getvalue(),'xlsx')

def test_excel_no_silent_cell_truncation(store):
    store.ingest([{**JOB,'description':'x'*40000}])
    with pytest.raises(ValueError,match='cell limit'):export_jobs(store,'xlsx')
    assert len(json.loads(export_jobs(store))[0]['description'])==40000

def test_no_clobber_atomic_export(tmp_path,monkeypatch):
    p=tmp_path/'export.json';publish_new(p,b'original')
    with pytest.raises(FileExistsError):publish_new(p,b'replace')
    assert p.read_bytes()==b'original' and not list(tmp_path.glob('*.tmp'))
    def fail(*args):raise OSError('fixture')
    monkeypatch.setattr('jobhunter_app.files.os.link',fail)
    with pytest.raises(OSError):publish_new(tmp_path/'new.json',b'new')
    assert not (tmp_path/'new.json').exists() and not list(tmp_path.glob('*.tmp'))

def test_uaa_database_projection_preserves_skip_and_determinism(store,tmp_path):
    # No document generation or downstream application process is invoked.
    evaluation={**JOB,'success':True,'global_score':4.5,'recommendation':'apply','evaluated_at':'2026-10-29T12:00:00Z'}
    d=legacy(tmp_path,[evaluation]);migrate(store,d,False)
    a=tmp_path/'a.jsonl';b=tmp_path/'b.jsonl'
    first=export_uaa(store,a);second=export_uaa(store,b)
    assert first['exported']==0 and first['skipped'][0]['reason']=='missing_documents'
    assert a.read_bytes()==b.read_bytes()==b''
    with pytest.raises(ValueError):export_uaa(store,a)


def test_history_export_not_limited_to_first_page(store,tmp_path):
    d=legacy(tmp_path,[{**JOB,'observation':i} for i in range(1001)]);migrate(store,d,False)
    assert len(json.loads(export_history(store))['historical_records'])==1001


def test_uaa_positive_projection_matches_pinned_exporter(store,tmp_path):
    from tests.test_export_queue_wq2 import make_eval
    from utils.queue_exporter import build_queue_entries, _detect_source
    evaluation=make_eval(tmp_path,external_job_id='1001',evaluated_at='2026-10-29T12:00:00Z')
    store.save_profile({**PROFILE,'original_cv':evaluation['cv_pdf_path']})
    d=legacy(tmp_path,[evaluation]);migrate(store,d,False)
    before={p.name:p.read_bytes() for p in tmp_path.glob('*.pdf')}
    a=tmp_path/'one.jsonl';b=tmp_path/'two.jsonl'
    result=export_uaa(store,a);export_uaa(store,b)
    job=store.jobs()[0]
    expected,_=build_queue_entries([evaluation],{evaluation['url']:{**job,'source':_detect_source(evaluation['url'])}}, {})
    assert result['exported']==1 and json.loads(a.read_text())==expected[0]
    assert a.read_bytes()==b.read_bytes() and {p.name:p.read_bytes() for p in tmp_path.glob('*.pdf')}==before


def test_migration_alias_cannot_overwrite_reviewed_identity(store,tmp_path):
    store.ingest([JOB]);jid=store.jobs()[0]['id'];alias='https://example.org/syndicated/1'
    with store.connect() as db:db.execute('INSERT INTO job_aliases VALUES(?,?,?)',(alias,jid,'fixture'))
    d=legacy(tmp_path,[{**JOB,'url':alias,'title':'Syndicated title','company':'Syndicator'}])
    migrate(store,d,False)
    assert store.job(jid)['title']==JOB['title'] and history(store)[0]['job_id']==jid


def test_import_state_rollback_on_history_error(store):
    with store.connect() as db:db.execute("CREATE TRIGGER fail BEFORE INSERT ON status_history BEGIN SELECT RAISE(ABORT,'fixture'); END")
    before=counts(store)
    with pytest.raises(sqlite3.DatabaseError):import_records(store,[{**JOB,'status':'saved'}],with_state=True)
    assert counts(store)==before


def test_simultaneous_migrations_do_not_duplicate_history(store,tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    d=legacy(tmp_path)
    # File lock is deliberately fail-fast; callers can explicitly rerun after completion.
    from filelock import Timeout
    def run():
        try:return migrate(store,d,False)
        except Timeout:return {'busy':True}
    with ThreadPoolExecutor(2) as pool:results=list(pool.map(lambda _:run(),range(2)))
    assert any('result' in x for x in results)
    assert counts(store)['historical_records']==1 and counts(store)['migration_batches']==1


def test_history_and_state_access_do_not_cross_profiles(store,tmp_path):
    store.save_profile(PROFILE,'private');d=legacy(tmp_path)
    migrate(store,d,False,'private');jid=store.jobs()[0]['id']
    assert history(store,'default',jid)==[] and len(history(store,'private',jid))==1
    assert export_uaa(store,tmp_path/'default.jsonl')['exported']==0


def test_excel_expansion_limit_rejected(monkeypatch):
    import jobhunter_app.spreadsheets as sheets
    monkeypatch.setattr(sheets,'MAX_EXPANDED',10)
    with pytest.raises(ValueError,match='expansion'):parse_jobs(xlsx([['a']],['url']),'xlsx')


def test_excel_dtd_entities_rejected():
    content=xlsx([[JOB['url'],JOB['title'],JOB['company']]],['url','title','company']);out=io.BytesIO()
    with ZipFile(io.BytesIO(content)) as src,ZipFile(out,'w',ZIP_DEFLATED) as z:
        for info in src.infolist():
            value=src.read(info.filename)
            if info.filename=='xl/worksheets/sheet1.xml':value=b'<!DOCTYPE x [<!ENTITY e "inert">]>'+value
            z.writestr(info,value)
    with pytest.raises(ValueError,match='Invalid XLSX'):parse_jobs(out.getvalue(),'xlsx')


def test_excel_missing_optional_dependency_explained(monkeypatch):
    import builtins
    original=builtins.__import__
    def without(name,*a,**k):
        if name=='openpyxl':raise ImportError('fixture')
        return original(name,*a,**k)
    content=xlsx([[JOB['url'],JOB['title'],JOB['company']]],['url','title','company'])
    monkeypatch.setattr(builtins,'__import__',without)
    with pytest.raises(ValueError,match='optional'):parse_jobs(content,'xlsx')


def test_uaa_never_assumes_legacy_tailored_cv_is_original(store,tmp_path):
    from tests.test_export_queue_wq2 import make_eval
    evaluation=make_eval(tmp_path)
    d=legacy(tmp_path,[evaluation]);migrate(store,d,False)
    r=export_uaa(store,tmp_path/'queue.jsonl')
    assert r['exported']==0 and r['skipped'][0]['reason']=='original_cv_not_selected'


def test_uaa_uses_selected_original_not_historical_cv(store,tmp_path):
    from tests.test_export_queue_wq2 import make_eval
    evaluation=make_eval(tmp_path)
    original=tmp_path/'candidate-original.pdf';original.write_bytes(b'%PDF-1.4 immutable synthetic fixture')
    store.save_profile({**PROFILE,'original_cv':str(original)})
    d=legacy(tmp_path,[evaluation]);migrate(store,d,False)
    target=tmp_path/'queue.jsonl';r=export_uaa(store,target)
    assert r['exported']==1 and json.loads(target.read_text())['cv_pdf']==str(original)
    assert r['original_cv_sha256']==sha256(original.read_bytes()).hexdigest()
    assert original.read_bytes()==b'%PDF-1.4 immutable synthetic fixture'


def test_uaa_timestamp_comparison_normalizes_timezones(store,tmp_path):
    from tests.test_export_queue_wq2 import make_eval
    evaluation=make_eval(tmp_path)
    store.save_profile({**PROFILE,'original_cv':evaluation['cv_pdf_path']})
    old={**evaluation,'evaluated_at':'2026-10-29T15:00:00+03:00','global_score':4.0}
    new={**evaluation,'evaluated_at':'2026-10-29T13:00:00+00:00','global_score':4.5}
    d=legacy(tmp_path,[old,new]);migrate(store,d,False)
    target=tmp_path/'queue.jsonl';export_uaa(store,target)
    assert json.loads(target.read_text())['score']==4.5


def test_uaa_ambiguous_equal_timestamp_is_not_arbitrarily_selected(store,tmp_path):
    from tests.test_export_queue_wq2 import make_eval
    evaluation=make_eval(tmp_path,evaluated_at='2026-10-29T12:00:00+00:00')
    store.save_profile({**PROFILE,'original_cv':evaluation['cv_pdf_path']})
    d=legacy(tmp_path,[evaluation,{**evaluation,'global_score':4.0}]);migrate(store,d,False)
    result=export_uaa(store,tmp_path/'queue.jsonl')
    assert result['exported']==0 and result['skipped'][0]['reason']=='ambiguous_latest_historical_evaluation'


def test_uaa_naive_timestamp_is_preserved_but_not_guessed(store,tmp_path):
    d=legacy(tmp_path,[{**JOB,'evaluated_at':'2026-10-29T12:00:00','success':True}]);migrate(store,d,False)
    result=export_uaa(store,tmp_path/'queue.jsonl')
    assert result['exported']==0 and result['skipped'][0]['reason']=='invalid_or_naive_evaluation_timestamp'
    assert history(store)[0]['payload']['evaluated_at']=='2026-10-29T12:00:00'


@pytest.mark.parametrize('content',[b'[1e309]',b'['*1500+b'0'+b']'*1500],ids=['overflow','nesting'])
def test_json_overflow_and_excessive_nesting_fail_safely(content):
    with pytest.raises(ValueError):parse_jobs(content,'json')


def test_oversized_csv_cell_is_a_safe_validation_error():
    with pytest.raises(ValueError):parse_jobs(b'url,title,company\na,'+b'x'*150000+b',c','csv')


def test_excel_unrepresentable_control_fails_without_losing_json_fact(store):
    store.ingest([{**JOB,'description':'original\x01text'}])
    with pytest.raises(ValueError,match='control character'):export_jobs(store,'xlsx')
    assert json.loads(export_jobs(store))[0]['description']=='original\x01text'


def test_excel_false_cached_dimensions_do_not_hide_cells():
    content=xlsx([[JOB['url'],JOB['title'],JOB['company']]],['url','title','company']);out=io.BytesIO()
    with ZipFile(io.BytesIO(content)) as src,ZipFile(out,'w',ZIP_DEFLATED) as z:
        for info in src.infolist():
            value=src.read(info.filename)
            if info.filename=='xl/worksheets/sheet1.xml':value=value.replace(b'ref="A1:C2"',b'ref="A1:A1"')
            z.writestr(info,value)
    assert parse_jobs(out.getvalue(),'xlsx')[0]['company']==JOB['company']


def test_csv_unterminated_quotation_is_not_silently_accepted():
    with pytest.raises(ValueError):parse_jobs(b'url,title,company\na,b,"company','csv')
