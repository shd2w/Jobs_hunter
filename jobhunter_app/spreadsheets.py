"""Optional, bounded values-only XLSX boundary. Never execute formulas or links."""
from __future__ import annotations
import io
import json
import re
from zipfile import ZipFile, BadZipFile
from defusedxml import ElementTree as ET
from defusedxml.common import DefusedXmlException

MAX_EXPANDED = 32 * 1024 * 1024
MAX_ROWS = 10000
FIELDS = ['id', 'url', 'title', 'company', 'location', 'description', 'status',
          'notes', 'deadline', 'first_seen', 'last_seen', 'attribution', 'source_notice', 'post_origins']


def _engine():
    try:
        import openpyxl
        return openpyxl
    except ImportError as exc:
        raise ValueError('Excel support is optional: install requirements-local-excel.txt') from exc


def read_xlsx(content: bytes, legacy=False) -> list[dict]:
    # Inspect compressed structure before openpyxl allocates worksheet objects.
    actual_rows, actual_columns = 0, 0
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 1000 or len({i.filename for i in infos}) != len(infos) or sum(i.file_size for i in infos) > MAX_EXPANDED:
                raise ValueError('XLSX archive exceeds safe expansion limits')
            for entry in infos:
                name = entry.filename.lower()
                if 'vbaproject' in name or 'externallinks/' in name or entry.flag_bits & 1:
                    raise ValueError('Macros, encrypted workbooks and external links are not supported')
                if name.endswith('.rels'):
                    relationships = ET.fromstring(archive.read(entry))
                    if any(x.attrib.get('TargetMode') == 'External' for x in relationships):
                        raise ValueError('External workbook relationships are not supported')
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    sheet = ET.fromstring(archive.read(entry))
                    if any(x.tag.rsplit('}', 1)[-1] == 'f' for x in sheet.iter()):
                        raise ValueError('Formula cells are not supported; supply a values-only copy')
                    for cell in sheet.iter():
                        if cell.tag.rsplit('}', 1)[-1] != 'c':
                            continue
                        coordinate = re.fullmatch(r'([A-Z]+)([1-9][0-9]*)', cell.attrib.get('r', ''))
                        if not coordinate:
                            raise ValueError('Invalid Excel cell coordinate')
                        column = 0
                        for letter in coordinate[1]:
                            column = column * 26 + ord(letter) - 64
                        actual_columns = max(actual_columns, column)
                        actual_rows = max(actual_rows, int(coordinate[2]))
                        if actual_columns > 40 or actual_rows > MAX_ROWS + 1:
                            raise ValueError('Excel actual cell bounds exceed limits')
                    rows = [x for x in sheet.iter() if x.tag.rsplit('}', 1)[-1] == 'row']
                    if len(rows) > MAX_ROWS + 1 or any(int(x.attrib.get('r', 0)) > MAX_ROWS + 1 for x in rows):
                        raise ValueError('Excel row limit is 10,000 records')
    except (BadZipFile, ET.ParseError, DefusedXmlException, KeyError) as exc:
        raise ValueError('Invalid XLSX container') from exc
    engine = _engine()
    try:
        workbook = engine.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError('Invalid XLSX workbook') from exc
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('Supply exactly one data worksheet; no sheet is silently ignored')
        sheet = workbook.worksheets[0]
        if sheet.sheet_state != 'visible' or (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 40:
            raise ValueError('Hidden or oversized worksheets are not supported')
        # Read actual bounded XML cells, not attacker-supplied cached dimensions.
        iterator = sheet.iter_rows(max_row=actual_rows or 1, max_col=actual_columns or 1)
        header_cells = next(iterator, [])
        headers = [x.value for x in header_cells]
        if not headers or any(not isinstance(h, str) or not h.strip() for h in headers) or len(set(headers)) != len(headers):
            raise ValueError('Excel requires unique nonempty text headers')
        aliases = {'Company': 'company', 'Title': 'title', 'URL': 'url', 'Date Posted': 'date_posted'} if legacy else {}
        out = []
        for cells in iterator:
            if all(c.value is None for c in cells):
                continue
            row = {}
            for header, cell in zip(headers, cells):
                if cell.data_type in ('f', 'e'):
                    raise ValueError('Formula/error cells are not supported')
                value = cell.value
                # Store exact typed legacy data; canonical intake is strings-only.
                if hasattr(value, 'isoformat'):
                    value = value.isoformat()
                if legacy:
                    row[header] = value
                else:
                    if value is not None and not isinstance(value, str):
                        raise ValueError('Canonical Excel job fields must be text')
                    row[header] = value or ''
            if legacy:
                row = {**row, **{v: row.get(k) or '' for k, v in aliases.items()}}
            out.append(row)
        return out
    finally:
        workbook.close()


def write_xlsx(rows: list[dict]) -> bytes:
    if len(rows) > MAX_ROWS:
        raise ValueError('Excel export exceeds 10,000 rows; use JSON or CSV')
    engine = _engine()
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book = engine.Workbook()
    sheet = book.active
    sheet.title = 'Opportunities'
    sheet.freeze_panes = 'A2'
    sheet.append(FIELDS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='245A81')
    for row in rows:
        values = []
        for field in FIELDS:
            value = row.get(field, '')
            if not isinstance(value, str):
                value = json.dumps(value, ensure_ascii=False, allow_nan=False)
            if any((ord(c) < 32 and c not in '\t\n\r') or ord(c) in (0xFFFE, 0xFFFF) or 0xD800 <= ord(c) <= 0xDFFF for c in value):
                raise ValueError('Excel cannot represent a stored control character; use JSON')
            if len(value) > 32767:
                raise ValueError('Excel cell limit exceeded; use JSON or CSV to avoid truncation')
            values.append(value)
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'  # A literal '=' remains text, without altering the fact.
            cell.number_format = '@'
            cell.alignment = Alignment(vertical='top')
    for i, field in enumerate(FIELDS, 1):
        sheet.column_dimensions[get_column_letter(i)].width = 50 if field in ('description', 'notes') else 26
    sheet.auto_filter.ref = sheet.dimensions
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()
